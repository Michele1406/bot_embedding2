import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

wb = openpyxl.Workbook()

font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
border_thin = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

# 1. Foglio Ricette
ws_ricette = wb.active
ws_ricette.title = "Ricette"
ws_ricette.append([
    "id_ricetta",
    "nome_piatto",
    "categoria",
    "stile_cucina",
    "canali_adatti",
    "canali_sconsigliati",
    "descrizione_breve",
    "note_composizione",
])

esempi_ricette = [
    [
        "PRIMO_MARE_01",
        "Fusillone con Alici sott'olio, Mollica croccante e Colatura",
        "primo",
        "mare",
        "ristorante, bistrot, trattoria",
        "bar",
        "Un primo piatto di carattere che esalta il sapore autentico del mare unito alla consistenza della pasta trafilata a bronzo.",
        "Mantecare a fuoco spento con olio EVO e colatura; finire con mollica tostata.",
    ],
    [
        "TAGLIERE_PUGLIA_01",
        "Tagliere Tradizione Murgiana",
        "tagliere",
        "pugliese",
        "ristorante, bottega, pub, wine_bar",
        "",
        "Selezione curata di eccellenze casearie e norcineria artigianale accompagnata da taralli e conserve.",
        "Step 1: proporre 6 salumi -> Step 2: 6 formaggi -> Step 3: taralli, olive e sottoli.",
    ],
    [
        "TRIS_APERITIVO_01",
        "Tris Aperitivo Sfizioso da Banco",
        "tris_bar",
        "aperitivo",
        "bar, pub, wine_bar, bistrot",
        "",
        "Tris secco pronto da servire per accompagnare cocktail e bollicine senza bisogno di cucina.",
        "Tre ciotoline: tarallini artigianali, olive giganti da tavola e frutta secca tostata (no pane carasau).",
    ],
    [
        "SECONDO_TERRA_01",
        "Hamburger Gourmet con Caciotta e Peperone Crusco",
        "secondo",
        "terra",
        "pub, bistrot, hamburgeria, ristorante",
        "bar",
        "Hamburger di maiale nero valorizzato da una fonduta leggera di caciotta artigianale e croccantezza vegetale.",
        "Servire con contorno di patate o verdure sott'olio ben scolate.",
    ],
]
for r in esempi_ricette:
    ws_ricette.append(r)

# 2. Foglio Ingredienti
ws_ingr = wb.create_sheet(title="Ingredienti")
ws_ingr.append([
    "id_ricetta",
    "nome_ingrediente_generico",
    "categoria_attesa",
    "ruolo",
    "quantita_indicativa",
    "note_ingrediente",
])

esempi_ingr = [
    [
        "PRIMO_MARE_01",
        "fusilli o pasta lunga trafilata",
        "Dispensa",
        "protagonista",
        "100g",
        "Grano duro italiano trafilato a bronzo",
    ],
    [
        "PRIMO_MARE_01",
        "alici in salamoia o sott'olio",
        "Mare",
        "protagonista",
        "40g",
        "Alici saporite deliscate",
    ],
    [
        "PRIMO_MARE_01",
        "olio extravergine di oliva",
        "Dispensa",
        "secondario",
        "20g",
        "Olio fruttato medio",
    ],
    [
        "PRIMO_MARE_01",
        "taralli o pane da tostare per mollica",
        "Dispensa",
        "opzionale",
        "15g",
        "Sbriciolati finemente in padella",
    ],
    [
        "TAGLIERE_PUGLIA_01",
        "capocollo o salume stagionato",
        "Salumi",
        "protagonista",
        "1 tagliere",
        "Salume artigianale a fette sottili",
    ],
    [
        "TAGLIERE_PUGLIA_01",
        "formaggio stagionato a pasta dura",
        "Formaggi",
        "protagonista",
        "1 tagliere",
        "Vaccino o caprino saporito",
    ],
    [
        "TAGLIERE_PUGLIA_01",
        "taralli classici",
        "Dispensa",
        "secondario",
        "1 ciotolina",
        "Accompagnamento croccante",
    ],
    [
        "TAGLIERE_PUGLIA_01",
        "olive da tavola",
        "Dispensa",
        "opzionale",
        "1 ciotolina",
        "Bella di Cerignola o simili",
    ],
    [
        "TRIS_APERITIVO_01",
        "tarallini pugliesi all'olio",
        "Dispensa",
        "protagonista",
        "1 ciotolina",
        "Formato snack",
    ],
    [
        "TRIS_APERITIVO_01",
        "olive verdi giganti",
        "Dispensa",
        "protagonista",
        "1 ciotolina",
        "Senza condimento piccante",
    ],
    [
        "TRIS_APERITIVO_01",
        "mandorle o frutta secca tostata",
        "Dispensa",
        "protagonista",
        "1 ciotolina",
        "Salata o naturale",
    ],
    [
        "SECONDO_TERRA_01",
        "hamburger di maiale nero o manzo",
        "Carne",
        "protagonista",
        "200g",
        "Carne fresca o gelo da cuocere",
    ],
    [
        "SECONDO_TERRA_01",
        "caciotta o formaggio fondente",
        "Formaggi",
        "secondario",
        "40g",
        "Sciolto in cottura",
    ],
    [
        "SECONDO_TERRA_01",
        "sott'oli o peperoni",
        "Dispensa",
        "opzionale",
        "30g",
        "Per contrasto agrodolce o sapido",
    ],
]
for r in esempi_ingr:
    ws_ingr.append(r)

for sheet in [ws_ricette, ws_ingr]:
    for cell in sheet[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        sheet.column_dimensions[col_letter].width = min(
            max(max_len + 3, 14), 45
        )

wb.save("ricettario.xlsx")