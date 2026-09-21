#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script per esportare la tassonomia assegnata a ciascun prodotto in SO FOOD.
Legge i prodotti direttamente dal database vettoriale (ChromaDB) di 'bot embedding 2'
e genera un file Excel (.xlsx) ordinato, formattato e completo di riepilogo.

Colonne principali richieste:
- Codice Fornitore
- Codice Prodotto
- Tassonomia (Reparto, Categoria Tassonomia, Sottocategoria, Percorso Completo)
"""

import os
import re
import sys
from pathlib import Path
import pandas as pd
import chromadb
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configurazione percorsi
CARTELLA_BASE = Path(__file__).resolve().parent
PERCORSO_DB = CARTELLA_BASE / "database_vettoriale"
NOME_COLLEZIONE = "catalogo_sofood"
NOME_FILE_OUTPUT = CARTELLA_BASE / "Tassonomia_Prodotti_SoFood.xlsx"


def pulisci_nome_prodotto(testo_doc: str) -> str:
    """Estrae e pulisce il nome del prodotto dalla prima riga del documento."""
    if not testo_doc:
        return ""
    # Rimuove BOM e caratteri spuri
    testo = testo_doc.replace("\ufeff", "").strip()
    righe = [r.strip() for r in testo.split("\n") if r.strip()]
    if not righe:
        return ""
    prima_riga = righe[0]
    # Rimuove prefissi come 'PRODOTTO: ' o simili
    nome_pulito = re.sub(r"^[^\w]*PRODOTTO:\s*", "", prima_riga, flags=re.IGNORECASE).strip()
    return nome_pulito


def esporta_tassonomia_excel(output_path: Path = NOME_FILE_OUTPUT):
    print("=" * 70)
    print("ESPORTAZIONE TASSONOMIA PRODOTTI SO FOOD -> EXCEL")
    print("=" * 70)
    print(f"[1/4] Connessione a ChromaDB ({PERCORSO_DB})...")

    if not PERCORSO_DB.exists():
        print(f"[ERRORE] Il percorso database vettoriale '{PERCORSO_DB}' non esiste.")
        return False

    client = chromadb.PersistentClient(path=str(PERCORSO_DB))
    collezione = client.get_collection(NOME_COLLEZIONE)
    totale_records = collezione.count()
    print(f"[OK] Collezione '{NOME_COLLEZIONE}' trovata: {totale_records} record totali.")

    print("[2/4] Lettura metadati e documenti da ChromaDB...")
    dati = collezione.get(limit=totale_records, include=["metadatas", "documents"])
    metadatas = dati["metadatas"]
    documents = dati["documents"]

    lista_prodotti = []
    lista_info_logistica = []

    for m, doc in zip(metadatas, documents):
        cod_prod = m.get("codice_prodotto")
        cod_forn = m.get("codice_fornitore")
        nome_forn = m.get("nome_fornitore", "")
        reparto = m.get("reparto", "")
        cat_tass = m.get("categoria_tassonomia", "")
        sotto = m.get("sottocategoria", "")
        cat_prod = m.get("categoria_prodotto", "")
        nome_prod = pulisci_nome_prodotto(doc)

        # Costruisce la tassonomia gerarchica completa
        pezzi_tassonomia = [p for p in [reparto, cat_tass, sotto] if p]
        tassonomia_completa = " > ".join(pezzi_tassonomia) if pezzi_tassonomia else "Non assegnata"

        riga = {
            "Codice Fornitore": cod_forn if cod_forn else "",
            "Nome Fornitore": nome_forn,
            "Codice Prodotto": cod_prod if cod_prod else "",
            "Nome Prodotto": nome_prod,
            "Reparto": reparto,
            "Categoria Tassonomia": cat_tass,
            "Sottocategoria": sotto,
            "Tassonomia Completa": tassonomia_completa,
            "Categoria Prodotto": cat_prod,
        }

        # Se manca codice_prodotto o è una scheda logistica/calendario
        if not cod_prod or cat_prod == "INFO LOGISTICA E ORDINI":
            riga["Descrizione"] = nome_prod if nome_prod else doc.split("\n")[0][:100]
            lista_info_logistica.append(riga)
        else:
            lista_prodotti.append(riga)

    df_prodotti = pd.DataFrame(lista_prodotti)
    # Ordinamento pulito per Fornitore e Codice Prodotto
    df_prodotti.sort_values(
        by=["Nome Fornitore", "Codice Fornitore", "Codice Prodotto"],
        inplace=True,
        ignore_index=True
    )

    print(f"[OK] Estratti {len(df_prodotti)} prodotti effettivi e {len(lista_info_logistica)} schede calendario/logistica.")

    # Creazione Foglio di Riepilogo Tassonomia
    df_riepilogo = (
        df_prodotti.groupby(["Reparto", "Categoria Tassonomia", "Sottocategoria"])
        .size()
        .reset_index(name="Numero Prodotti")
        .sort_values(by=["Reparto", "Categoria Tassonomia", "Numero Prodotti"], ascending=[True, True, False])
    )

    print(f"[3/4] Generazione file Excel ({output_path.name})...")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Foglio 1: Prodotti con Tassonomia
        df_prodotti.to_excel(writer, sheet_name="Tassonomia Prodotti", index=False)
        # Foglio 2: Riepilogo aggregato
        df_riepilogo.to_excel(writer, sheet_name="Riepilogo Tassonomie", index=False)
        # Foglio 3: Schede logistiche / calendari
        if lista_info_logistica:
            df_logistica = pd.DataFrame(lista_info_logistica)
            df_logistica.to_excel(writer, sheet_name="Info Logistica e Calendari", index=False)

        # Formattazione grafica con openpyxl
        wb = writer.book

        # Stili
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Blu scuro elegante
        fill_sub_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        font_data = Font(name="Calibri", size=10)
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )
        fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

        # Formatta ogni foglio
        for ws in wb.worksheets:
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "A2"  # Blocca la prima riga

            header_fill = fill_header if ws.title == "Tassonomia Prodotti" else fill_sub_header

            # Header styling
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = font_header
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

            # Data rows styling
            for row_idx in range(2, ws.max_row + 1):
                is_even = (row_idx % 2 == 0)
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = font_data
                    cell.border = border_thin
                    if is_even:
                        cell.fill = fill_zebra
                    # Allineamento
                    header_name = str(ws.cell(row=1, column=col_idx).value or "")
                    if "Codice" in header_name or "Numero" in header_name:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-fit larghezza colonne
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                # Applica larghezza con margine (max 65 per evitare colonne abnormi)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

            # Abilita filtri automatici
            ws.auto_filter.ref = ws.dimensions

    print("[4/4] Formattazione e salvataggio completati!")
    print("=" * 70)
    print(f"File generato con successo: {output_path.resolve()}")
    print(f"  - Prodotti catalogati: {len(df_prodotti)}")
    print(f"  - Categorie uniche:    {len(df_prodotti['Categoria Tassonomia'].unique())}")
    print(f"  - Sottocategorie:      {len(df_prodotti['Sottocategoria'].unique())}")
    print(f"  - Reparti:             {len(df_prodotti['Reparto'].unique())}")
    print("=" * 70)
    return True


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    esporta_tassonomia_excel()
